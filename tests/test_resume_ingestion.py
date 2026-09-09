from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from src.infrastructure.llm_client import FakeModelClient
from src.services.resume_ingestion_service import (
    ResumeEvidenceExtractor,
    ResumeIngestionError,
    extract_resume_text,
)


def _evidence_output(*, claim: str, skills: list[str]) -> dict[str, object]:
    return {
        "evidence": [
            {
                "type": "project",
                "title": "Agent 项目",
                "claim": claim,
                "skills": skills,
            }
        ]
    }


@pytest.mark.asyncio
async def test_extractor_retries_ungrounded_output_and_allows_pdf_line_wraps() -> None:
    client = FakeModelClient(
        outputs=[
            _evidence_output(claim="使用了不存在的 Kubernetes", skills=["Kubernetes"]),
            _evidence_output(
                claim="使用 Python 开发 Agent 服务",
                skills=["Python", "Agent"],
            ),
        ]
    )
    extractor = ResumeEvidenceExtractor(client, validation_retries=1)

    evidence = await extractor.extract("项目经历\n使用 Python 开发\nAgent 服务")

    assert len(client.requests) == 2
    assert evidence[0].claim == "使用 Python 开发 Agent 服务"
    assert evidence[0].skills == ["Python", "Agent"]
    assert "上次输出未通过校验" in client.requests[1].messages[-1].content


@pytest.mark.asyncio
async def test_extractor_rejects_api_facts_missing_from_resume() -> None:
    extractor = ResumeEvidenceExtractor(
        FakeModelClient(
            output=_evidence_output(
                claim="负责推荐系统研发",
                skills=["推荐系统"],
            )
        ),
        validation_retries=0,
    )

    with pytest.raises(ResumeIngestionError) as caught:
        await extractor.extract("负责 Agent 应用研发")

    assert caught.value.code == "resume_evidence_ungrounded"


def test_extract_resume_text_reads_docx_paragraphs_and_tables() -> None:
    target = BytesIO()
    document = Document()
    document.add_paragraph("Python 项目经历")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Agent 工作流"
    document.save(target)

    text = extract_resume_text(
        content=target.getvalue(),
        media_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
    )

    assert text == "Python 项目经历\nAgent 工作流"
