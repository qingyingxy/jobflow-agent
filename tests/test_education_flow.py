from __future__ import annotations

from copy import deepcopy
from io import BytesIO

import httpx
import pytest
from docx import Document
from sqlalchemy import select

from src.domain.education import (
    EducationCondition,
    EducationRecord,
    validate_condition_grounding,
)
from src.domain.eligibility import (
    CandidateProfileInput,
    EligibilityInput,
    SearchPreferences,
)
from src.domain.job import RawJobDocument, StructuredJobDescription
from src.domain.models import EvidenceItem, UserProfile
from src.domain.runs import JobParseResult
from src.infrastructure.llm_client import FakeModelClient
from src.main import app
from src.services.education_evaluator import evaluate_education
from src.services.eligibility_checker import check_eligibility
from src.services.jd_analysis_service import JDAnalysisService
from src.services.local_evidence_matcher import match_requirement_locally
from src.services.product_jd_education_parser import ProductJDParser
from src.services.profile_service import ProfileService
from src.services.resume_ingestion_service import (
    ResumeEvidenceExtractor,
    ResumeIngestionError,
    ResumeIngestionService,
    is_current_resume_asset,
)

CLAUSE = "本科要求计算机相关专业，硕士不限专业"
RAW = "招聘开发工程师，任职要求：" + CLAUSE + "。欢迎投递。"
CONDITION = {
    "relation": "any_of",
    "branches": [
        {
            "source_text": "本科要求计算机相关专业",
            "degree": {"operator": "eq", "level": "bachelor", "source_text": "本科"},
            "major": {
                "mode": "related",
                "values": ["计算机"],
                "source_text": "计算机相关专业",
            },
        },
        {
            "source_text": "硕士不限专业",
            "degree": {"operator": "eq", "level": "master", "source_text": "硕士"},
            "major": {"mode": "any", "source_text": "不限专业"},
        },
    ],
}


def record(level="bachelor", major="机械专业", year=None):
    label = {"bachelor": "本科", "master": "硕士", "unknown": "学历未确认"}[level]
    return EducationRecord(
        id=f"edu_{level}",
        level=level,
        major_raw=major,
        graduation_year=year,
        source_text=f"{label}是{major or '专业未填'}",
    )


def model_outputs(condition=None):
    return [
        {
            "facts": {},
            "requirements": [{"source_text": CLAUSE, "level": "required"}],
            "responsibilities": [],
        },
        {
            "decisions": [
                {
                    "requirement_index": 0,
                    "relation": "any_of",
                    "items": ["本科要求计算机相关专业", "硕士不限专业"],
                    "reason": "不同学历的备选资格",
                    "education": condition or CONDITION,
                }
            ]
        },
    ]


@pytest.mark.parametrize(
    ("records", "complete", "expected"),
    [
        ([record(), record("master", "金融专业")], False, "pass"),
        ([record("bachelor", "计算机科学与技术")], True, "pass"),
        ([record()], True, "unknown"),
        ([record()], False, "unknown"),
        ([record("bachelor", "交叉学科")], True, "unknown"),
        ([record("unknown", "计算机科学")], True, "unknown"),
        ([record("master", None)], False, "pass"),
        ([], False, "unknown"),
    ],
)
def test_qualification_three_valued(records, complete, expected):
    result = evaluate_education(
        EducationCondition.model_validate(CONDITION),
        CandidateProfileInput(educations=records, education_history_complete=complete),
    )
    assert result.result == expected
    if expected == "pass" and any(r.level == "master" for r in records):
        assert result.evidence[0].id == "edu_master"


@pytest.mark.parametrize(
    ("records", "complete", "expected"),
    [
        ([record(), record("master", "金融")], True, "unknown"),
        ([record("bachelor", "计算机科学")], True, "fail"),
        ([record("bachelor", "计算机科学"), record("master", "金融")], True, "pass"),
        ([record("bachelor", "计算机科学")], False, "unknown"),
    ],
)
def test_both_education_branches_must_pass_for_all_of(records, complete, expected):
    condition = EducationCondition(relation="all_of", branches=CONDITION["branches"])
    result = evaluate_education(
        condition,
        CandidateProfileInput(educations=records, education_history_complete=complete),
    )
    assert result.result == expected


def test_major_only_condition_does_not_invent_degree_requirement():
    condition = EducationCondition(
        relation="all_of",
        branches=[
            {
                "source_text": "计算机相关专业",
                "degree": {"operator": "not_specified"},
                "major": CONDITION["branches"][0]["major"],
            }
        ],
    )
    validate_condition_grounding(condition, "计算机相关专业")
    profile = CandidateProfileInput(educations=[record("unknown", "软件工程")])
    assert evaluate_education(condition, profile).result == "pass"


def test_same_record_scopes_major_and_graduation_year():
    branch = deepcopy(CONDITION["branches"][0])
    branch["major"] = {
        "mode": "exact",
        "values": ["计算机科学"],
        "source_text": "计算机科学",
    }
    condition = EducationCondition(relation="all_of", branches=[branch])
    profile = CandidateProfileInput(
        educations=[record(), record("master", "计算机科学", 2027)],
        education_history_complete=True,
    )
    assert evaluate_education(condition, profile).result == "fail"
    branch.update(graduation_years=[2027], graduation_source_text="2027届")
    branch["major"] = {"mode": "not_specified"}
    condition = EducationCondition(relation="all_of", branches=[branch])
    profile.educations[0].graduation_year = 2024
    assert evaluate_education(condition, profile).result == "fail"


def test_global_graduation_year_cannot_be_borrowed_from_other_degree():
    from src.domain.job import FieldEvidence, JobRequirement

    source = "本科要求计算机相关专业"
    condition = EducationCondition(
        relation="all_of",
        branches=[
            {
                **CONDITION["branches"][0],
                "scope": "application",
                "scope_source_text": "面向2027届毕业生，本科要求计算机相关专业",
            }
        ],
    )
    requirement = JobRequirement(
        category="education",
        name=source,
        description=source,
        education=condition,
        evidence=[FieldEvidence(field_path="requirements[0]", source_text=source)],
    )
    job = StructuredJobDescription(
        requirements=[requirement],
        graduation_years=[2027],
        field_evidence=[
            FieldEvidence(field_path="graduation_years", source_text="面向2027届毕业生")
        ],
    )
    profile = CandidateProfileInput(
        educations=[
            record("bachelor", "计算机科学", 2024),
            record("master", "金融", 2027),
        ],
        graduation_year=2027,
        education_history_complete=True,
    )
    result = check_eligibility(
        EligibilityInput(profile=profile, job=job, preferences=SearchPreferences()),
        include_search_preferences=False,
    )
    assert result.eligible == "fail"
    match = match_requirement_locally(
        user_id="a",
        requirement=requirement,
        evidence=[],
        profile=profile,
        education_checks=result.checks,
    )
    assert match.support_level == "unsupported"


def test_unknown_status_is_not_graduated_and_unresolved_is_not_pass():
    branch = deepcopy(CONDITION["branches"][1])
    profile = CandidateProfileInput(
        educations=[record("master", "金融")], education_history_complete=True
    )
    branch["status"] = "graduated"
    assert (
        evaluate_education(
            EducationCondition(relation="all_of", branches=[branch]), profile
        ).result
        == "unknown"
    )
    branch.update(status="any", unresolved=True)
    assert (
        evaluate_education(
            EducationCondition(relation="all_of", branches=[branch]), profile
        ).result
        == "unknown"
    )


def test_no_major_constraint_cannot_remove_degree_constraint():
    condition = deepcopy(CONDITION)
    condition["branches"][1]["degree"] = {"operator": "any", "source_text": "不限专业"}
    with pytest.raises(ValueError, match="专业不限不等于学历不限"):
        validate_condition_grounding(
            EducationCondition.model_validate(condition), CLAUSE
        )


@pytest.mark.asyncio
async def test_runtime_two_calls_compile_predicate_and_local_match():
    client = FakeModelClient(outputs=model_outputs())
    parsed = await ProductJDParser(client).parse(RawJobDocument(raw_content=RAW))
    requirement = parsed.structured_jd.requirements[0]
    assert len(client.requests) == 2
    assert "EducationCondition" in client.requests[1].json_schema["$defs"]
    assert parsed.parser_version == "product-jd-parser-v6"
    assert requirement.category == "education"
    assert requirement.education.branches[1].degree.level == "master"
    assert requirement.evidence[0].source_text == CLAUSE
    assert parsed.structured_jd.required_skills is None
    profile = CandidateProfileInput(educations=[record(), record("master", "金融")])
    match = match_requirement_locally(
        user_id="a", requirement=requirement, evidence=[], profile=profile
    )
    assert match.support_level == "supported"
    assert match.evidence_ids == []
    assert match.education_evidence[0].id == "edu_master"
    # Isolate the education condition from the existing unspecified-cohort gate.
    job = StructuredJobDescription(
        **{**parsed.structured_jd.model_dump(), "graduation_years": []}
    )
    eligibility = check_eligibility(
        EligibilityInput(profile=profile, job=job, preferences=SearchPreferences()),
        include_search_preferences=False,
    )
    assert eligibility.eligible == "pass"
    assert {c.rule_name for c in eligibility.checks} == {
        "graduation_year",
        "education_qualification_0",
    }


@pytest.mark.asyncio
async def test_invalid_runtime_predicate_is_repaired_in_second_stage():
    bad = deepcopy(CONDITION)
    bad["branches"][1]["degree"] = {"operator": "any", "source_text": "不限专业"}
    outputs = model_outputs(bad) + [model_outputs()[1]]
    client = FakeModelClient(outputs=outputs)
    parsed = await ProductJDParser(client).parse(RawJobDocument(raw_content=RAW))
    assert len(client.requests) == 3
    assert (
        parsed.structured_jd.requirements[0].education.branches[1].degree.operator
        == "eq"
    )


@pytest.mark.asyncio
async def test_resume_extract_save_replace_and_manual_preservation(db_session):
    source = "本科 机械专业 2020-2024\n硕士 金融专业 2024-2027 预计毕业"
    output = {
        "evidence": [
            {
                "type": "education",
                "title": "本科",
                "claim": source.splitlines()[0],
                "skills": [],
                "education": {
                    "level": "bachelor",
                    "major_raw": "机械专业",
                    "enrollment_year": 2020,
                    "enrollment_source_text": "2020-2024",
                    "graduation_year": 2024,
                    "graduation_source_text": "2020-2024",
                },
            },
            {
                "type": "education",
                "title": "硕士",
                "claim": source.splitlines()[1],
                "skills": [],
                "education": {
                    "level": "master",
                    "major_raw": "金融专业",
                    "enrollment_year": 2024,
                    "enrollment_source_text": "2024-2027",
                    "graduation_year": 2027,
                    "graduation_source_text": "2024-2027 预计毕业",
                    "graduation_year_kind": "expected",
                    "status": "enrolled",
                },
            },
        ]
    }
    document = Document()
    document.add_paragraph(source)
    content = BytesIO()
    document.save(content)
    client = FakeModelClient(output=output)
    service = ResumeIngestionService(db_session, ResumeEvidenceExtractor(client))
    await service.ingest(
        user_id="a",
        asset_id="asset1",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        content=content.getvalue(),
    )
    profile = db_session.get(UserProfile, "a")
    assert [e["graduation_year"] for e in profile.educations] == [2024, 2027]
    assert [e["enrollment_year"] for e in profile.educations] == [2020, 2024]
    assert profile.educations[1]["graduation_year_kind"] == "expected"
    assert (
        profile.degree == "硕士"
        and profile.major == "金融专业"
        and profile.graduation_year == 2027
    )
    assert not profile.education_history_complete
    assert (
        list(db_session.scalars(select(EvidenceItem))) == []
    )  # no second editable copy
    assert is_current_resume_asset(db_session, user_id="a", asset_id="asset1")
    assert not is_current_resume_asset(db_session, user_id="b", asset_id="asset1")
    manual = record("bachelor", "数学").model_dump(mode="json")
    ProfileService(db_session).upsert(
        "a",
        {
            "educations": [*profile.educations, manual],
            "education_history_complete": True,
        },
    )
    extracted = await service.extractor.extract(source)
    service._replace_extracted_evidence(
        user_id="a", asset_id="asset2", extracted=extracted
    )
    assert len(profile.educations) == 3
    assert profile.educations[0]["source"] == "manual"
    assert profile.education_history_complete  # unchanged history needs no repeat review
    assert is_current_resume_asset(db_session, user_id="a", asset_id="asset2")


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "changes",
    [
        {"major_raw": "金融专业"},
        {"level": "master"},
        {"status": "graduated"},
        {"graduation_year": 2027, "graduation_source_text": "2027"},
    ],
)
async def test_resume_education_cannot_borrow_other_record_facts(changes):
    output = {
        "evidence": [
            {
                "type": "education",
                "title": "本科",
                "claim": "本科 机械专业 2024",
                "skills": [],
                "education": {"level": "bachelor", "major_raw": "机械专业", **changes},
            }
        ]
    }
    extractor = ResumeEvidenceExtractor(
        FakeModelClient(output=output), validation_retries=0
    )
    with pytest.raises(ResumeIngestionError, match="教育|学历|毕业"):
        await extractor.extract("本科 机械专业 2024\n硕士 金融专业 2027 已毕业")


@pytest.mark.asyncio
async def test_profile_api_persists_education_and_invalidates_analysis(db_session):
    headers = {"X-User-ID": "a"}
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.put(
            "/api/profile",
            headers=headers,
            json={
                "educations": [
                    record().model_dump(),
                    record("master", "金融", 2027).model_dump(),
                ],
                "education_history_complete": True,
            },
        )
        assert response.status_code == 200, response.text
        assert response.json()["degree"] == "硕士"
        assert (
            await client.get("/api/profile", headers={"X-User-ID": "b"})
        ).status_code == 404
        # Build a posting through the normal import API, then use a controlled model.
        imported = await client.post(
            "/api/jobs/import-text", headers=headers, json={"raw_content": RAW}
        )
        assert imported.status_code in {200, 201}, imported.text
        job_id = imported.json()["id"]
        parser = ProductJDParser(FakeModelClient(outputs=model_outputs()))
        execution = await JDAnalysisService(db_session, parser=parser).analyze(
            user_id="a", job_id=job_id
        )
        assert execution.analysis.matches[0]["support_level"] == "supported"
        checks = execution.analysis.eligibility["checks"]
        assert next(c for c in checks if c["field"] == "educations")["result"] == "pass"
        assert (
            next(c for c in checks if c["field"] == "educations")["education_evidence"][
                0
            ]["id"]
            == "edu_master"
        )
        assert db_session.scalar(select(JobParseResult)).structured_jd["requirements"][
            0
        ]["education"]
        reject_legacy = await client.put(
            "/api/profile", headers=headers, json={"major": "计算机"}
        )
        assert reject_legacy.status_code == 422
        update = await client.put(
            "/api/profile",
            headers=headers,
            json={
                "educations": [record().model_dump()],
                "education_history_complete": True,
            },
        )
        assert update.status_code == 200
        assert JDAnalysisService(db_session).latest(user_id="a", job_id=job_id) is None
        rerun = await JDAnalysisService(db_session, parser=parser).analyze(
            user_id="a", job_id=job_id
        )
        assert rerun.analysis.eligibility["eligible"] == "unknown"
        assert rerun.analysis.matches[0]["support_level"] == "needs_confirmation"
